"""Document loading and processing utilities."""
import os
import csv
from typing import List, Dict, Callable
from pathlib import Path
import pypdf
from docx import Document


class DocumentLoader:
    """Load and process documents from various formats."""
    
    def __init__(self, documents_path: str):
        """Initialize the document loader.
        
        Args:
            documents_path: Path to the directory containing documents
        """
        self.documents_path = Path(documents_path)
        
        # Register file type handlers (easily expandable)
        self._loaders: Dict[str, Callable[[Path], str]] = {
            '.txt': self.load_txt,
            '.md': self.load_txt,
            '.pdf': self.load_pdf,
            '.docx': self.load_docx,
            '.csv': self.load_csv,
            '.xlsx': self.load_excel,
        }
    
    @property
    def supported_extensions(self) -> set:
        """Get set of supported file extensions."""
        return set(self._loaders.keys())
    
    def load_txt(self, file_path: Path) -> str:
        """Load text from a .txt or .md file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def load_pdf(self, file_path: Path) -> str:
        """Load text from a PDF file."""
        text = []
        with open(file_path, 'rb') as f:
            pdf_reader = pypdf.PdfReader(f)
            for page in pdf_reader.pages:
                text.append(page.extract_text())
        return '\n'.join(text)
    
    def load_docx(self, file_path: Path) -> str:
        """Load text from a Word document."""
        doc = Document(file_path)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)
        return '\n'.join(text)
    
    def load_csv(self, file_path: Path) -> str:
        """Load text from a CSV file.
        
        Converts CSV data to a readable text format with headers and rows.
        """
        text = []
        with open(file_path, 'r', encoding='utf-8', newline='') as f:
            csv_reader = csv.reader(f)
            rows = list(csv_reader)
            
            if not rows:
                return ""
            
            # First row is typically headers
            if rows:
                headers = rows[0]
                text.append("Headers: " + ", ".join(headers))
                text.append("-" * 50)
                
                # Process data rows
                for i, row in enumerate(rows[1:], 1):
                    row_text = []
                    for header, value in zip(headers, row):
                        row_text.append(f"{header}: {value}")
                    text.append(f"Row {i}: " + " | ".join(row_text))
        
        return '\n'.join(text)
    
    def load_excel(self, file_path: Path) -> str:
        """Load text from an Excel (.xlsx) file.
        
        Converts Excel data to a readable text format with sheet names, headers, and rows.
        """
        from openpyxl import load_workbook
        
        text = []
        workbook = load_workbook(filename=file_path, read_only=True)
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text.append(f"Sheet: {sheet_name}")
            rows = list(sheet.iter_rows(values_only=True))
            
            if not rows:
                continue
            
            # First row is typically headers
            headers = rows[0]
            text.append("Headers: " + ", ".join([str(h) for h in headers if h is not None]))
            text.append("-" * 50)
            
            # Process data rows
            for i, row in enumerate(rows[1:], 1):
                row_text = []
                for header, value in zip(headers, row):
                    if header is not None:
                        row_text.append(f"{header}: {value}")
                text.append(f"Row {i}: " + " | ".join(row_text))
        
        return '\n'.join(text)
    
    def register_loader(self, extension: str, loader_func: Callable[[Path], str]):
        """Register a custom loader for a new file type.
        
        This allows easy expansion to support additional file formats.
        
        Args:
            extension: File extension (e.g., '.json', '.xml')
            loader_func: Function that takes a Path and returns string content
            
        Example:
            def load_json(file_path: Path) -> str:
                import json
                with open(file_path, 'r') as f:
                    data = json.load(f)
                return json.dumps(data, indent=2)
            
            loader.register_loader('.json', load_json)
        """
        if not extension.startswith('.'):
            extension = '.' + extension
        self._loaders[extension.lower()] = loader_func
        print(f"Registered loader for {extension} files")
    
    def load_document(self, file_path: Path) -> Dict[str, str]:
        """Load a single document and return its content with metadata.
        
        Args:
            file_path: Path to the document
            
        Returns:
            Dictionary with 'content', 'filename', and 'filepath'
        """
        extension = file_path.suffix.lower()
        
        # Use registered loader if available
        if extension in self._loaders:
            loader_func = self._loaders[extension]
            content = loader_func(file_path)
        else:
            raise ValueError(
                f"Unsupported file format: {extension}. "
                f"Supported formats: {', '.join(sorted(self.supported_extensions))}"
            )
        
        return {
            'content': content,
            'filename': file_path.name,
            'filepath': str(file_path)
        }
    
    def load_all_documents(self) -> List[Dict[str, str]]:
        """Load all supported documents from the documents directory.
        
        Returns:
            List of dictionaries containing document content and metadata
        """
        if not self.documents_path.exists():
            os.makedirs(self.documents_path, exist_ok=True)
            print(f"Created documents directory: {self.documents_path}")
            return []
        
        documents = []
        for file_path in self.documents_path.rglob('*'):
            if file_path.is_file() and file_path.suffix.lower() in self.supported_extensions:
                try:
                    doc = self.load_document(file_path)
                    documents.append(doc)
                    print(f"Loaded: {file_path.name}")
                except Exception as e:
                    print(f"Error loading {file_path.name}: {str(e)}")
        
        return documents
    
    def chunk_text(self, text: str, chunk_size: int = 1000, 
                   chunk_overlap: int = 200) -> List[str]:
        """Split text into overlapping chunks.
        
        Args:
            text: Text to split
            chunk_size: Maximum size of each chunk
            chunk_overlap: Number of characters to overlap between chunks
            
        Returns:
            List of text chunks
        """
        chunks = []
        start = 0
        text_length = len(text)
        
        while start < text_length:
            end = start + chunk_size
            chunk = text[start:end]
            chunks.append(chunk)
            start += chunk_size - chunk_overlap
        
        return chunks
